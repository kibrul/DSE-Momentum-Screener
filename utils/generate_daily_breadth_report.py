"""
generate_daily_breadth_report.py

Builds a day-by-day DSE market breadth summary for the last ~3 months and
saves it to an Excel file. Reuses compute_live_breadth() from utils/live.py
UNCHANGED — the same function the app's Live tab uses for the current
moment — applied here once per historical trading day instead of once for
"right now".

compute_live_breadth() expects a snapshot DataFrame with a "% Change"
column (that's what it actually reads: up_4pct, down_4pct, advancers,
decliners are all derived from that one column). So for each historical
day, this script reconstructs that same shape from Close-to-Close % change
across the whole loaded universe, then hands it to compute_live_breadth()
exactly as if it were a live snapshot for that day.

Run:
    python generate_daily_breadth_report.py

Output:
    dse_daily_breadth_summary.xlsx  (columns: Date, Up 4%+, Down 4%+, Advancers, Decliners)

Requires network access to dsebd.org (via bdshare) — run this locally,
not in a restricted sandbox.
"""
import os
import sys
from datetime import date, timedelta

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import pandas as pd
from openpyxl.styles import Font, Alignment

from utils.data import fetch_full_market_history, get_all_trading_codes, filter_equity_codes
from utils.live import compute_live_breadth

# ---------------- Configuration ----------------
LOOKBACK_DAYS = 92          # ~3 calendar months; DSE trades ~5 days/week so this yields ~60 trading days
EQUITIES_ONLY = True        # exclude mutual funds/bonds/etc, same heuristic filter as the main app
OUTPUT_PATH = "dse_daily_breadth_summary.xlsx"


def build_daily_breadth_table(price_data: dict[str, pd.DataFrame]) -> pd.DataFrame:
    """
    Reconstructs a day-by-day breadth summary by calling compute_live_breadth()
    once per historical trading day, exactly as the Live tab calls it for
    the current moment.
    """
    # Wide matrix: rows = date, columns = symbol, values = Close.
    # Building a DataFrame from a dict of Series aligns everything on the
    # union of all dates automatically (outer join), so stocks that didn't
    # trade on a given day just show up as NaN there.
    close_matrix = pd.DataFrame({symbol: df["Close"] for symbol, df in price_data.items()})
    close_matrix = close_matrix.sort_index()

    # Day-over-day % change per symbol, vectorized — row t compares close_t to close_{t-1}.
    pct_change_matrix = close_matrix.pct_change() * 100

    rows = []
    for dt, row in pct_change_matrix.iterrows():
        changes = row.dropna()
        if changes.empty:
            continue  # first day in the window (no prior close) or a non-trading day

        # Reconstruct the exact shape compute_live_breadth() expects.
        snapshot = pd.DataFrame({"% Change": changes})
        stats = compute_live_breadth(snapshot)

        rows.append({
            "Date": dt.date() if hasattr(dt, "date") else dt,
            "Up 4%+": stats["up_4pct"],
            "Down 4%+": stats["down_4pct"],
            "Advancers": stats["advancers"],
            "Decliners": stats["decliners"],
        })

    return pd.DataFrame(rows).sort_values("Date").reset_index(drop=True)


def save_to_excel(df: pd.DataFrame, path: str) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Daily Breadth")
        ws = writer.sheets["Daily Breadth"]

        header_font = Font(name="Arial", bold=True)
        body_font = Font(name="Arial")

        for cell in ws[1]:
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font

        widths = {"A": 14, "B": 10, "C": 11, "D": 12, "E": 12}
        for col, w in widths.items():
            ws.column_dimensions[col].width = w

        for cell in list(ws["A"])[1:]:
            cell.number_format = "yyyy-mm-dd"
            cell.alignment = Alignment(horizontal="center")
        for col in ["B", "C", "D", "E"]:
            for cell in list(ws[col])[1:]:
                cell.alignment = Alignment(horizontal="center")


def main():
    end_date = date.today()
    start_date = end_date - timedelta(days=LOOKBACK_DAYS)
    print(f"Fetching DSE full-market history: {start_date.isoformat()} to {end_date.isoformat()} ...")

    price_data = fetch_full_market_history(start_date.isoformat(), end_date.isoformat())
    if not price_data:
        print("No data returned: DSE could not be reached from this machine and there is no saved snapshot (see the message above).")
        return
    print(f"Loaded history for {len(price_data)} symbols.")

    if EQUITIES_ONLY:
        all_codes = get_all_trading_codes() or sorted(price_data)   # [] when DSE's code list is unreachable
        equity_codes = set(filter_equity_codes(all_codes))
        price_data = {k: v for k, v in price_data.items() if k in equity_codes}
        print(f"Filtered to {len(price_data)} equity symbols (excluding mutual funds/bonds/etc — heuristic).")

    print("Computing day-by-day breadth summary via compute_live_breadth()...")
    daily_df = build_daily_breadth_table(price_data)

    if daily_df.empty:
        print("No daily breadth data could be computed — check the fetched history isn't empty.")
        return

    save_to_excel(daily_df, OUTPUT_PATH)
    print(f"Saved {len(daily_df)} trading days to {OUTPUT_PATH}")
    print(daily_df.tail(5).to_string(index=False))


if __name__ == "__main__":
    main()
